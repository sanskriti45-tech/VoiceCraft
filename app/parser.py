"""Extract plain text from an uploaded resume / portfolio file."""

import io
import re

MAX_BYTES = 5 * 1024 * 1024  # 5 MB upload limit
MAX_CHARS = 20000  # text sent to the LLM
ALLOWED = {".pdf", ".docx", ".txt", ".md"}


class ParseError(Exception):
    """The file could not be read; message is safe to show the user."""


def _ext(filename: str) -> str:
    name = (filename or "").lower()
    return name[name.rfind("."):] if "." in name else ""


def _pdf_text(data: bytes) -> str:
    from pypdf import PdfReader

    try:
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ParseError("This PDF is password-protected. Remove the password and try again.")
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except ParseError:
        raise
    except Exception:
        raise ParseError("Could not read this PDF. Try exporting it again or upload a DOCX.")


def _docx_text(data: bytes) -> str:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception:
        raise ParseError("Could not read this DOCX file. Is it a valid Word document?")
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(filename: str, data: bytes) -> str:
    ext = _ext(filename)
    if ext not in ALLOWED:
        raise ParseError("Unsupported file type. Upload a PDF, DOCX, TXT or MD file.")

    if ext == ".pdf":
        text = _pdf_text(data)
    elif ext == ".docx":
        text = _docx_text(data)
    else:
        text = data.decode("utf-8", errors="replace")

    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()

    if len(text) < 50:
        raise ParseError(
            "Could not find readable text in this file. If it is a scanned PDF, "
            "upload a text-based PDF or a DOCX instead."
        )
    return text[:MAX_CHARS]
